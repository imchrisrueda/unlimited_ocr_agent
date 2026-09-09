import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.fieldnotes.render.excel import export_estadillo_xlsx, publish_estadillo_xlsx
from src.fieldnotes.schemas.evidence import EvidenceValue
from src.fieldnotes.schemas.estadillo import EstadilloDocument, EstadilloPage, EstadilloRow


class TestEstadilloExcelRoundTrip(unittest.TestCase):
    def _document(self) -> EstadilloDocument:
        row = EstadilloRow(
            source_page=1,
            col=EvidenceValue[int](normalized=1, source_page=1),
            fil=EvidenceValue[int](normalized=26, source_page=1),
            especie=EvidenceValue[str](raw="Ap", normalized="P", source_page=1),
            altura_cm=EvidenceValue[float](normalized=4.5, source_page=1),
            bbch=EvidenceValue[str](raw="22", normalized="22", source_page=1),
        )
        return EstadilloDocument(source_file="sample.pdf", pages=[EstadilloPage(page_number=1, rows=[row])])

    def test_export_and_publish_accepts_spanish_decimal(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            workbook_path = folder / "datos.xlsx"
            csv_path = folder / "datos.csv"
            export_estadillo_xlsx(self._document(), workbook_path)
            workbook = load_workbook(workbook_path)
            workbook["Datos"]["E2"] = "4,75"
            workbook.save(workbook_path)
            self.assertEqual(publish_estadillo_xlsx(workbook_path, csv_path), 1)
            with csv_path.open(encoding="utf-8", newline="") as stream:
                rows = list(csv.reader(stream))
            self.assertEqual(rows[0], ["id", "col", "fil", "especie", "altura_cm", "foto", "bbch", "observaciones"])
            self.assertEqual(rows[1][4], "4.75")

    def test_publish_rejects_unrecognized_species(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            workbook_path = folder / "datos.xlsx"
            export_estadillo_xlsx(self._document(), workbook_path)
            workbook = load_workbook(workbook_path)
            workbook["Datos"]["D2"] = "At"
            workbook.save(workbook_path)
            with self.assertRaisesRegex(ValueError, "especie debe ser"):
                publish_estadillo_xlsx(workbook_path, folder / "datos.csv")

    def test_publish_updates_document_json_with_reviewed_values(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            workbook_path = folder / "datos.xlsx"
            csv_path = folder / "datos.csv"
            document_path = folder / "document.json"
            document = self._document()
            document_path.write_text(document.model_dump_json(indent=2), encoding="utf-8")
            export_estadillo_xlsx(document, workbook_path)
            workbook = load_workbook(workbook_path)
            workbook["Datos"]["E2"] = "4,75"
            workbook.save(workbook_path)

            publish_estadillo_xlsx(workbook_path, csv_path, document_path)

            updated = EstadilloDocument.model_validate_json(document_path.read_text(encoding="utf-8"))
            self.assertEqual(updated.pages[0].rows[0].altura_cm.normalized, 4.75)
            self.assertEqual(updated.pages[0].rows[0].source_page, 1)


if __name__ == "__main__":
    unittest.main()
