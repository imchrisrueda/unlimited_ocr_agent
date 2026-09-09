"""Excel review workbook import/export for canonical estadillo CSV data."""
from __future__ import annotations

import csv
import io
import math
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.fieldnotes.schemas.estadillo import EstadilloDocument
from src.fieldnotes.schemas.evidence import EvidenceValue

CSV_HEADER = ("id", "col", "fil", "especie", "altura_cm", "foto", "bbch", "observaciones")
SHEET_NAME = "Datos"


def _value(evidence):
    if evidence is None:
        return None
    return evidence.normalized if evidence.normalized is not None else evidence.raw


def export_estadillo_xlsx(doc: EstadilloDocument, destination: Path) -> Path:
    """Export the canonical records as a typed, Excel-friendly review workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(list(CSV_HEADER))
    for page in doc.pages:
        for row in page.rows:
            sheet.append([
                _value(row.id), _value(row.col), _value(row.fil), _value(row.especie),
                _value(row.altura_cm), _value(row.foto), _value(row.bbch), _value(row.observaciones),
            ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 9
    sheet.column_dimensions["C"].width = 9
    sheet.column_dimensions["D"].width = 11
    sheet.column_dimensions["E"].width = 14
    sheet.column_dimensions["F"].width = 11
    sheet.column_dimensions["G"].width = 11
    sheet.column_dimensions["H"].width = 45
    for cell in sheet["E"][1:]:
        cell.number_format = "0.0#"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def _text(value, field: str, row_number: int) -> str:
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(f"Fila {row_number}: no se admiten formulas en '{field}'.")
    return str(value).strip()


def _integer(value, field: str, row_number: int) -> str:
    text = _text(value, field, row_number)
    if not text:
        return ""
    try:
        number = int(text)
    except ValueError as exc:
        raise ValueError(f"Fila {row_number}: '{field}' debe ser un entero.") from exc
    return str(number)


def _height(value, row_number: int) -> str:
    text = _text(value, "altura_cm", row_number).replace(",", ".")
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"Fila {row_number}: 'altura_cm' debe ser numerica.") from exc
    if not math.isfinite(number):
        raise ValueError(f"Fila {row_number}: 'altura_cm' debe ser finita.")
    return format(number, "g")


def _atomic_write_text(destination: Path, content: str) -> None:
    """Write text atomically without leaving an interrupted publication behind."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=f".tmp_{destination.name}_", dir=destination.parent)
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
        os.replace(temporary_name, destination)
    except Exception:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)
        raise


def _replace_row_values(document: EstadilloDocument, records: list[list[str]]) -> EstadilloDocument:
    """Apply reviewed table cells to their original rows while retaining page provenance."""
    reviewed = document.model_copy(deep=True)
    rows = [row for page in reviewed.pages for row in page.rows]
    if len(rows) != len(records):
        raise ValueError(
            "El Excel debe conservar exactamente el mismo número y orden de registros que document.json."
        )

    def text_evidence(value: str, page: int):
        return EvidenceValue[str](raw=value, normalized=value, source_page=page) if value else None

    def int_evidence(value: str, page: int):
        return EvidenceValue[int](raw=value, normalized=int(value), source_page=page) if value else None

    def float_evidence(value: str, page: int):
        return EvidenceValue[float](raw=value, normalized=float(value), source_page=page) if value else None

    for row, values in zip(rows, records):
        row.id = text_evidence(values[0], row.source_page)
        row.col = int_evidence(values[1], row.source_page)
        row.fil = int_evidence(values[2], row.source_page)
        row.especie = text_evidence(values[3], row.source_page)
        row.altura_cm = float_evidence(values[4], row.source_page)
        row.foto = text_evidence(values[5], row.source_page)
        row.bbch = text_evidence(values[6], row.source_page)
        row.observaciones = text_evidence(values[7], row.source_page)
    return reviewed


def publish_estadillo_xlsx(
    workbook_path: Path,
    csv_path: Path,
    document_json_path: Path | None = None,
) -> int:
    """Validate a review workbook and publish its CSV and optional document JSON together."""
    if not workbook_path.is_file():
        raise FileNotFoundError(f"No existe el libro de revision: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)

    def invalid(message: str) -> None:
        workbook.close()
        raise ValueError(message)
    if SHEET_NAME not in workbook.sheetnames:
        invalid(f"El libro debe contener la hoja '{SHEET_NAME}'.")
    sheet = workbook[SHEET_NAME]
    header = tuple(_text(cell.value, "cabecera", 1) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if header != CSV_HEADER:
        invalid(f"Cabecera invalida. Se esperaba: {','.join(CSV_HEADER)}")

    records: list[list[str]] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(CSV_HEADER)), start=2):
        values = [cell.value for cell in cells]
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        species = _text(values[3], "especie", row_number).upper()
        if species and species not in {"P", "H", "R", "M"}:
            invalid(f"Fila {row_number}: especie debe ser P, H, R o M.")
        bbch = _text(values[6], "bbch", row_number)
        if bbch and (len(bbch) != 2 or not bbch.isascii() or not bbch.isdigit()):
            invalid(f"Fila {row_number}: bbch debe tener dos digitos ASCII.")
        records.append([
            _text(values[0], "id", row_number), _integer(values[1], "col", row_number),
            _integer(values[2], "fil", row_number), species, _height(values[4], row_number),
            _text(values[5], "foto", row_number), bbch, _text(values[7], "observaciones", row_number),
        ])
    workbook.close()

    csv_output = io.StringIO(newline="")
    writer = csv.writer(csv_output, lineterminator="\n")
    writer.writerow(CSV_HEADER)
    writer.writerows(records)

    updated_document: EstadilloDocument | None = None
    if document_json_path is not None:
        if not document_json_path.is_file():
            raise FileNotFoundError(f"No existe document.json: {document_json_path}")
        original_document = EstadilloDocument.model_validate_json(document_json_path.read_text(encoding="utf-8"))
        updated_document = _replace_row_values(original_document, records)

    if updated_document is not None:
        _atomic_write_text(document_json_path, updated_document.model_dump_json(indent=2))
    _atomic_write_text(csv_path, csv_output.getvalue())
    return len(records)
